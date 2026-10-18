import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Tuple

# Styling palettes
NAVY_HEADER = "1E293B"
LIGHT_NAVY = "334155"
ACCENT_BLUE = "2563EB"
BG_LIGHT_GRAY = "F8FAFC"
BORDER_GRAY = "E2E8F0"
TEXT_WHITE = "FFFFFF"

def create_component_template_excel(course_info: Dict[str, Any], component: Dict[str, Any], students: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks Entry"

    # Header styling
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    meta_font = Font(name="Calibri", size=10, color="475569")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color=TEXT_WHITE)
    border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )

    ws["A1"] = f"Symbiosis Institute of Technology (SIT), Nagpur | {course_info['code']} - {course_info['title']}"
    ws["A1"].font = title_font
    ws["A2"] = f"Component: {component['name']} ({component['code']}) | Maximum Marks: {component['max_marks']} | Department: {course_info.get('department', 'CSE')} | Semester: {course_info['semester']}"
    ws["A2"].font = meta_font
    ws["A3"] = "Instructions: Enter marks between 0 and maximum marks. For absent students, type 'AB' or 'Absent'."
    ws["A3"].font = Font(name="Calibri", size=9, italic=True, color="64748B")

    headers = ["Roll Number", "Student Name", "Email", f"Marks Obtained (Max: {component['max_marks']})", "Status (Present/Absent)", "Comments/Remarks"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, s in enumerate(students, 6):
        ws.cell(row=r_idx, column=1, value=s.get("roll_number", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=2, value=s.get("name", ""))
        ws.cell(row=r_idx, column=3, value=s.get("email", ""))
        marks_cell = ws.cell(row=r_idx, column=4, value="")
        marks_cell.alignment = Alignment(horizontal="right")
        ws.cell(row=r_idx, column=5, value="Present").alignment = Alignment(horizontal="center")
        ws.cell(row=r_idx, column=6, value="")

        for c_idx in range(1, 7):
            ws.cell(row=r_idx, column=c_idx).border = border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def parse_component_upload(file_bytes: bytes, filename: str, max_marks: float) -> Tuple[List[Dict[str, Any]], List[str]]:
    parsed_entries = []
    errors = []

    if filename.endswith(".csv"):
        text = file_bytes.decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        # Find header row
        header_row_idx = 0
        for i, row in enumerate(rows):
            row_str = " ".join(row).lower()
            if "roll" in row_str or "student" in row_str:
                header_row_idx = i
                break
        
        header = [h.strip().lower() for h in rows[header_row_idx]]
        roll_idx = next((i for i, h in enumerate(header) if "roll" in h), 0)
        marks_idx = next((i for i, h in enumerate(header) if "mark" in h or "obtained" in h or "score" in h), -1)
        if marks_idx == -1 and len(header) >= 4:
            marks_idx = 3 # fallback to 4th column

        status_idx = next((i for i, h in enumerate(header) if "status" in h or "absent" in h), -1)
        comment_idx = next((i for i, h in enumerate(header) if "comment" in h or "remark" in h), -1)

        for line_no, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
            if not row or not any(row):
                continue
            roll = row[roll_idx].strip() if len(row) > roll_idx else ""
            if not roll:
                continue

            marks_str = row[marks_idx].strip() if marks_idx != -1 and len(row) > marks_idx else ""
            status_str = row[status_idx].strip().lower() if status_idx != -1 and len(row) > status_idx else ""
            comment = row[comment_idx].strip() if comment_idx != -1 and len(row) > comment_idx else ""

            is_absent = (status_str in ("ab", "absent", "a")) or (marks_str.lower() in ("ab", "absent", "a"))
            val = None

            if not is_absent and marks_str != "":
                try:
                    val = float(marks_str)
                    if val < 0:
                        errors.append(f"Row {line_no} ({roll}): Marks cannot be negative ({val}).")
                        continue
                    if val > max_marks:
                        errors.append(f"Row {line_no} ({roll}): Marks ({val}) exceed maximum permissible ({max_marks}).")
                        continue
                except ValueError:
                    errors.append(f"Row {line_no} ({roll}): Invalid marks value '{marks_str}'.")
                    continue

            parsed_entries.append({
                "roll_number": roll,
                "marks_obtained": val if not is_absent else 0.0,
                "is_absent": is_absent,
                "comments": comment
            })

    else:
        # Excel .xlsx parsing
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        header_row_idx = 0
        for i, row in enumerate(rows):
            row_str = " ".join(str(cell or "") for cell in row).lower()
            if "roll" in row_str:
                header_row_idx = i
                break
        
        header = [str(h or "").strip().lower() for h in rows[header_row_idx]]
        roll_idx = next((i for i, h in enumerate(header) if "roll" in h), 0)
        marks_idx = next((i for i, h in enumerate(header) if "mark" in h or "score" in h or "obtained" in h), -1)
        if marks_idx == -1 and len(header) >= 4:
            marks_idx = 3

        status_idx = next((i for i, h in enumerate(header) if "status" in h or "absent" in h), -1)
        comment_idx = next((i for i, h in enumerate(header) if "comment" in h or "remark" in h), -1)

        for line_no, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
            if not row or not any(row):
                continue
            roll = str(row[roll_idx] or "").strip()
            if not roll:
                continue

            marks_val = row[marks_idx] if marks_idx != -1 and len(row) > marks_idx else None
            status_val = str(row[status_idx] or "").strip().lower() if status_idx != -1 and len(row) > status_idx else ""
            comment = str(row[comment_idx] or "").strip() if comment_idx != -1 and len(row) > comment_idx else ""

            is_absent = (status_val in ("ab", "absent", "a")) or (str(marks_val or "").strip().lower() in ("ab", "absent", "a"))
            val = None

            if not is_absent and marks_val is not None and str(marks_val).strip() != "":
                try:
                    val = float(marks_val)
                    if val < 0:
                        errors.append(f"Row {line_no} ({roll}): Marks cannot be negative ({val}).")
                        continue
                    if val > max_marks:
                        errors.append(f"Row {line_no} ({roll}): Marks ({val}) exceed maximum permissible ({max_marks}).")
                        continue
                except (ValueError, TypeError):
                    errors.append(f"Row {line_no} ({roll}): Invalid marks value '{marks_val}'.")
                    continue

            parsed_entries.append({
                "roll_number": roll,
                "marks_obtained": val if not is_absent else 0.0,
                "is_absent": is_absent,
                "comments": comment
            })

    return parsed_entries, errors

def generate_consolidated_master_excel(consolidated_data: Dict[str, Any], audit_logs: List[Dict[str, Any]] = None) -> bytes:
    wb = openpyxl.Workbook()
    
    # Sheet 1: Master Consolidated Sheet
    ws = wb.active
    ws.title = "Consolidated Internal Marks"

    # Theme colors
    navy_dark = "0F172A"
    blue_header = "1E40AF"
    blue_light = "DBEAFE"
    gray_light = "F1F5F9"
    border_style = Side(style='thin', color="CBD5E1")
    cell_border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    
    # Title Block
    ws.merge_cells("A1:M1")
    ws["A1"] = f"Symbiosis Institute of Technology (SIT), Nagpur | {consolidated_data['course_code']} - {consolidated_data['course_title']}"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A2:M2")
    instructor_str = f" | Instructor: {consolidated_data['instructor_name']}" if consolidated_data.get('instructor_name') else ""
    ws["A2"] = f"Department: {consolidated_data['department']} | Semester: {consolidated_data['semester']}{instructor_str} | Status: {consolidated_data['submission_status'].upper()} {'(LOCKED)' if consolidated_data['is_locked'] else ''}"
    ws["A2"].font = Font(name="Calibri", size=10, color="475569")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    # Analytics banner in sheet
    an = consolidated_data.get("analytics", {})
    ws.merge_cells("A3:M3")
    ws["A3"] = f"Summary: Enrolled: {an.get('total_students', 0)} | Class Avg: {an.get('mean_score', 0)}% | Median: {an.get('median_score', 0)}% | Pass Rate: {an.get('pass_percentage', 0)}% | Highest: {an.get('highest_score', 0)} | Lowest: {an.get('lowest_score', 0)}"
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="1E3A8A")

    # Header Row
    headers = ["Rank", "Roll No", "Student Name", "Email"]
    
    # Inspect sample student to get groups and components
    sample_student = consolidated_data["students"][0] if consolidated_data["students"] else None
    
    if sample_student:
        for grp in sample_student["groups"]:
            for comp in grp["components"]:
                headers.append(f"{comp['code']}\n(Max: {comp['max_marks']})")
            headers.append(f"{grp['group_name']}\n({grp['group_weightage']}%)")
    
    headers.extend(["Total Internal\n(100)", "Grade", "GP", "Result"])

    header_row_idx = 5
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=blue_header, end_color=blue_header, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    # Student Rows
    for r_idx, s in enumerate(consolidated_data["students"], start=header_row_idx + 1):
        row_vals = [
            s.get("rank", ""),
            s.get("roll_number", ""),
            s.get("student_name", ""),
            s.get("email", "")
        ]

        for grp in s["groups"]:
            for comp in grp["components"]:
                if comp["is_absent"]:
                    row_vals.append("AB")
                elif comp["marks_obtained"] is not None:
                    included_mark = f"{comp['marks_obtained']}" + ("*" if not comp["is_included"] else "")
                    row_vals.append(included_mark)
                else:
                    row_vals.append("-")
            # Group score
            row_vals.append(f"{grp['weighted_contribution']:.1f}")

        row_vals.extend([
            f"{s['total_marks']:.2f}",
            s["grade"],
            s["grade_point"],
            "PASS" if s["is_passed"] else "FAIL"
        ])

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = cell_border
            
            # Alignments & conditional styling
            if c_idx in (1, 2):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in (3, 4):
                cell.alignment = Alignment(horizontal="left")
            else:
                cell.alignment = Alignment(horizontal="center")

            # Result coloring
            if c_idx == len(row_vals):
                if val == "PASS":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                else:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    # Footnote about Best-of-N
    last_row = header_row_idx + len(consolidated_data["students"]) + 2
    ws.cell(row=last_row, column=1, value="* Note: Component scores with asterisk (*) were dropped according to Best-of-N rules.").font = Font(italic=True, size=9, color="64748B")

    # Auto-adjust widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(cell.value or '')) for cell in col[:header_row_idx + 10])
        ws.column_dimensions[col_letter].width = max(min(max_l + 4, 30), 10)

    # Sheet 2: Grade Distribution
    ws_grades = wb.create_sheet(title="Grade Analytics")
    ws_grades["A1"] = "Grade Scale & Class Distribution"
    ws_grades["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    
    g_headers = ["Letter Grade", "Grade Point", "Description", "Min Percentage", "Count", "Percentage of Class"]
    for c_idx, h in enumerate(g_headers, 1):
        c = ws_grades.cell(row=3, column=c_idx, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.alignment = Alignment(horizontal="center")

    for r_idx, g in enumerate(an.get("grade_distribution", []), 4):
        ws_grades.cell(row=r_idx, column=1, value=g["grade"]).alignment = Alignment(horizontal="center")
        ws_grades.cell(row=r_idx, column=2, value=g["grade_point"]).alignment = Alignment(horizontal="center")
        ws_grades.cell(row=r_idx, column=3, value=g["description"])
        ws_grades.cell(row=r_idx, column=4, value=f">= {g['min_percentage']}%").alignment = Alignment(horizontal="center")
        ws_grades.cell(row=r_idx, column=5, value=g["count"]).alignment = Alignment(horizontal="right")
        ws_grades.cell(row=r_idx, column=6, value=f"{g['percentage']}%").alignment = Alignment(horizontal="right")
        for c in range(1, 7):
            ws_grades.cell(row=r_idx, column=c).border = cell_border

    # Sheet 3: Audit Trail Log
    if audit_logs:
        ws_audit = wb.create_sheet(title="Audit Trail Log")
        ws_audit["A1"] = f"Complete Edit Traceability Log - {consolidated_data['course_code']}"
        ws_audit["A1"].font = Font(size=14, bold=True)
        
        a_headers = ["Timestamp", "Student Roll", "Student Name", "Component", "Old Mark", "New Mark", "Changed By", "Role", "Reason", "Source"]
        for c_idx, h in enumerate(a_headers, 1):
            c = ws_audit.cell(row=3, column=c_idx, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
            c.alignment = Alignment(horizontal="center")

        for r_idx, log in enumerate(audit_logs, 4):
            ws_audit.cell(row=r_idx, column=1, value=log.get("timestamp", ""))
            ws_audit.cell(row=r_idx, column=2, value=log.get("student_roll", "")).alignment = Alignment(horizontal="center")
            ws_audit.cell(row=r_idx, column=3, value=log.get("student_name", ""))
            ws_audit.cell(row=r_idx, column=4, value=log.get("component_name", ""))
            ws_audit.cell(row=r_idx, column=5, value=str(log.get("old_value") if log.get("old_value") is not None else "-")).alignment = Alignment(horizontal="right")
            ws_audit.cell(row=r_idx, column=6, value=str(log.get("new_value") if log.get("new_value") is not None else "-")).alignment = Alignment(horizontal="right")
            ws_audit.cell(row=r_idx, column=7, value=log.get("changed_by_name", ""))
            ws_audit.cell(row=r_idx, column=8, value=log.get("changed_by_role", ""))
            ws_audit.cell(row=r_idx, column=9, value=log.get("reason", ""))
            ws_audit.cell(row=r_idx, column=10, value=log.get("source", ""))
            for c in range(1, 11):
                ws_audit.cell(row=r_idx, column=c).border = cell_border

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
